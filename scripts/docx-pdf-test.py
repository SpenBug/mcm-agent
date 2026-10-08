"""真跑：造一个 .docx → 转成 .pdf → 验证 PDF 真的有效。

用 app 自带的 Python 环境（和运行时一致）。
"""
import os
import shutil
import subprocess
import sys
import tempfile

APP_PY = r"C:\Users\92182\AppData\Roaming\阿一古数模\python-env\Scripts\python.exe"
SCRIPT = r"D:\数学建模项目\mcm-agent\resources\skills\mcm-tools\scripts\mcm_docx.py"

WS = tempfile.mkdtemp(prefix="pdftest-")
os.makedirs(os.path.join(WS, ".mcm-agent"), exist_ok=True)
DOCX = os.path.join(WS, "论文.docx")

# ① 用 python-docx 造一个有中文、有标题、有表格的文档
print("=== ① 造测试 .docx ===")
mk = subprocess.run([APP_PY, "-c", '''
import sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

doc = Document()
st = doc.styles["Normal"]
st.font.name = "宋体"
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

doc.add_heading("基于多目标优化的生产调度模型", level=1)
doc.add_paragraph("摘要：本文针对某工厂的生产调度问题，建立了多目标优化模型。")
doc.add_paragraph("关键词：多目标优化；遗传算法；生产调度")
doc.add_heading("一、问题重述", level=2)
doc.add_paragraph("该问题要求我们在满足产能约束的前提下，最小化总完工时间。")
t = doc.add_table(rows=3, cols=3)
t.style = "Table Grid"
for i, row in enumerate([["工序", "耗时(分钟)", "机器"], ["A", "12", "M1"], ["B", "8", "M2"]]):
    for j, v in enumerate(row):
        t.cell(i, j).text = v
doc.add_heading("二、模型建立", level=2)
doc.add_paragraph("目标函数为 min Z = w1*Cmax + w2*Tardiness。")
doc.save(sys.argv[1])
print("docx 已生成")
''' , DOCX], capture_output=True, text=True, encoding="utf-8", errors="replace")
print("  ", mk.stdout.strip() or mk.stderr.strip()[:200])
print("   文件:", DOCX, "存在:", os.path.exists(DOCX),
      ("%.1f KB" % (os.path.getsize(DOCX) / 1024)) if os.path.exists(DOCX) else "")

if not os.path.exists(DOCX):
    print("  ✗ docx 没造出来，中止")
    sys.exit(2)

# ② 转 PDF
print()
print("=== ② 转 PDF（mcm_docx.py pdf）===")
r = subprocess.run([APP_PY, SCRIPT, "pdf", DOCX],
                   capture_output=True, text=True, encoding="utf-8", errors="replace",
                   cwd=WS, timeout=240)
print("  退出码:", r.returncode)
for line in (r.stdout or "").splitlines():
    print("   ", line)
if r.stderr.strip():
    print("   stderr:", r.stderr.strip()[:400])

PDF = os.path.join(WS, "论文.pdf")
ok_exists = os.path.exists(PDF)
print()
print("=== ③ 验证 PDF ===")
print("   文件存在:", "✓" if ok_exists else "✗")
if ok_exists:
    size = os.path.getsize(PDF)
    print("   大小: %.1f KB" % (size / 1024))
    with open(PDF, "rb") as f:
        head = f.read(8)
        f.seek(-64, os.SEEK_END)
        tail = f.read()
    print("   文件头:", head[:5])
    is_pdf = head.startswith(b"%PDF-")
    print("   是合法 PDF:", "✓" if is_pdf else "✗")
    # 用 pypdf 读一下页数，证明内容真的在
    try:
        from pypdf import PdfReader
        rd = PdfReader(PDF)
        txt = (rd.pages[0].extract_text() or "")
        print("   页数:", len(rd.pages))
        print("   首页含标题:", "✓" if "生产调度" in txt else "✗  " + txt[:60].replace("\n", " "))
        print("   首页含正文:", "✓" if "多目标优化" in txt else "✗")
        ok_pdf = is_pdf and len(rd.pages) >= 1 and "生产调度" in txt
    except Exception as e:                                   # noqa: BLE001
        print("   pypdf 读取失败:", e)
        ok_pdf = is_pdf
else:
    ok_pdf = False

print()
print("=== " + ("✓ .docx → .pdf 全链路通过" if ok_pdf else "✗ 未通过") + " ===")
shutil.rmtree(WS, ignore_errors=True)
sys.exit(0 if ok_pdf else 1)
